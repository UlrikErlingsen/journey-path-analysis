from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
import pytest

from tracesignal.design import audit_event_log
from tracesignal.errors import DataProblem
from tracesignal.examples import make_demo_events, make_starter_template
from tracesignal.io import (
    MAX_TABLE_COLUMNS,
    MAX_UPLOAD_BYTES,
    build_evidence_workbook,
    dataframe_csv_bytes,
    read_table,
    safe_frame,
)


ROOT = Path(__file__).parents[1]


def test_demo_is_deterministic() -> None:
    pd.testing.assert_frame_equal(make_demo_events(), make_demo_events())


def test_demo_is_wholly_fictional_and_has_real_sequences(demo_events) -> None:
    assert demo_events["journey_id"].str.startswith("JOURNEY-").all()
    assert demo_events["customer_id"].str.startswith("CUSTOMER-").all()
    assert demo_events["journey_id"].nunique() == 720
    assert len(demo_events) > demo_events["journey_id"].nunique()


def test_committed_demo_matches_generator() -> None:
    committed = pd.read_csv(ROOT / "examples" / "tracesignal-fictional-event-log.csv")
    generated = make_demo_events()
    pd.testing.assert_frame_equal(committed, generated, check_dtype=False)


def test_starter_template_has_contract_columns() -> None:
    template = make_starter_template()
    required = {"journey_id", "timestamp", "touchpoint", "converted"}
    assert required.issubset(template.columns)


def test_csv_round_trip(demo_events) -> None:
    payload = dataframe_csv_bytes(demo_events.head(20))
    read = read_table("events.csv", payload)
    assert len(read) == 20
    assert list(read.columns) == list(demo_events.columns)


def test_xlsx_round_trip(demo_events) -> None:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        demo_events.head(20).to_excel(writer, sheet_name="events", index=False)
    read = read_table("events.xlsx", output.getvalue())
    assert len(read) == 20


def test_evidence_workbook_contains_auditable_sheets(validated, result) -> None:
    payload = build_evidence_workbook(
        metadata={"app": "Trace Signal", "analysis_type": "descriptive"},
        audit=audit_event_log(validated),
        result=result,
    )
    sheets = set(pd.ExcelFile(BytesIO(payload)).sheet_names)
    assert {
        "read_me",
        "analysis_config",
        "transitions",
        "positional_roles",
        "dropoff",
        "paths",
        "removal_sensitivity",
        "memory_diagnostic",
        "limitations",
    }.issubset(sheets)


HOSTILE_TOUCHPOINT = '=HYPERLINK("http://evil.example","click")'


def _hostile_frame() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "touchpoint": [HOSTILE_TOUCHPOINT, "+SUM(A1:A9)", "-2+3", "@cmd", "Safe page", "bell\x07control"],
            "=danger_header": [1, 2, 3, 4, 5, 6],
        }
    )


def test_csv_export_neutralizes_formula_cells_and_headers() -> None:
    payload = dataframe_csv_bytes(_hostile_frame()).decode("utf-8")
    lines = payload.splitlines()
    assert lines[0].startswith("touchpoint,'=danger_header")
    assert "'" + HOSTILE_TOUCHPOINT.replace('"', '""') in payload
    assert "\n=HYPERLINK" not in payload and not any(line.startswith("=") for line in lines)
    assert "'+SUM(A1:A9)" in payload
    assert "'-2+3" in payload
    assert "'@cmd" in payload
    assert "\x07" not in payload


def test_safe_frame_strips_control_characters_but_keeps_text() -> None:
    cleaned = safe_frame(_hostile_frame())
    assert cleaned["touchpoint"].iloc[5] == "bellcontrol"
    assert cleaned["touchpoint"].iloc[4] == "Safe page"


def test_workbook_export_neutralizes_hostile_touchpoint(demo_events, config) -> None:
    from tracesignal import analyze_journeys, validate_event_log

    events = demo_events.copy()
    events.loc[events["touchpoint"] == "Webinar", "touchpoint"] = HOSTILE_TOUCHPOINT
    validated = validate_event_log(events)
    result = analyze_journeys(validated, config)
    payload = build_evidence_workbook(
        metadata={"app": "Trace Signal", "note": "=2+2"},
        audit=audit_event_log(validated),
        result=result,
    )
    book = load_workbook(BytesIO(payload))
    neutralized_cells = 0
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                # No exported string cell may be a live formula.
                assert not cell.value.startswith(("=", "+", "@"))
                if cell.value.startswith("'" + HOSTILE_TOUCHPOINT):
                    neutralized_cells += 1
                    assert cell.data_type == "s"
    assert neutralized_cells > 0


def test_oversized_upload_is_refused() -> None:
    with pytest.raises(DataProblem, match="50 MB"):
        read_table("big.csv", b"0" * (MAX_UPLOAD_BYTES + 1))


def test_empty_upload_is_refused() -> None:
    with pytest.raises(DataProblem, match="empty"):
        read_table("empty.csv", b"")


def test_too_many_columns_are_refused() -> None:
    header = ",".join(f"c{i}" for i in range(MAX_TABLE_COLUMNS + 1))
    payload = (header + "\n" + ",".join("1" for _ in range(MAX_TABLE_COLUMNS + 1))).encode("utf-8")
    with pytest.raises(DataProblem, match="column safety limit"):
        read_table("wide.csv", payload)


def test_unknown_extension_is_refused() -> None:
    with pytest.raises(DataProblem, match="CSV or XLSX"):
        read_table("data.parquet", b"PAR1")
